"""Regression tests for defects observed in the pre-repair resume output."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

import pytest
from sqlalchemy.orm import Session

from apps.resume_generation.schemas import (
    BulletRef,
    EducationEntry,
    ExperienceEntry,
    ProjectEntry,
    ResumeDocument,
)
from apps.resume_generation.services.ats_validation_service import validate_artifacts
from apps.resume_generation.services.docx_render_service import (
    extract_docx_text,
    render_docx,
)
from apps.resume_generation.services.generation_preflight import (
    ELIGIBILITY_BLOCK_MSG,
    generation_preflight,
)
from apps.resume_generation.services.generation_service import generate_resume_for_job
from apps.resume_generation.services.resume_corpus_service import normalize_text
from apps.resume_generation.services.resume_fact_service import scan_resume_corpus
from apps.resume_generation.services.resume_strategy_service import (
    build_resume_strategy,
)
from apps.resume_generation.services.structured_document_service import (
    _is_near_duplicate,
    build_deduped_skill_lines,
    build_resume_document,
    validate_document_structure,
)
from packages.core.db import SessionLocal, get_engine
from packages.core.domain.models import (
    Base,
    Company,
    JobPosting,
    ResumeFact,
    ResumeGeneration,
    make_dedupe_key,
    make_description_hash,
    make_uuid,
)
from tests.resume_generation.conftest import (
    build_conflicting_gpa_corpus,
    build_standard_corpus,
    write_resume_docx,
)


@pytest.fixture()
def session() -> Session:
    engine = get_engine()
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _job(
    session: Session,
    *,
    title: str,
    description: str,
    location: str = "United States",
    url: str = "https://example.com/apply",
) -> JobPosting:
    company = Company(
        canonical_name=f"Co-{make_uuid()[:8]}",
        normalized_key=make_uuid()[:8],
    )
    session.add(company)
    session.flush()
    job = JobPosting(
        id=make_uuid(),
        source="test",
        title=title,
        company_id=company.id,
        location=location,
        description_text=description,
        description_hash=make_description_hash(description),
        dedupe_key=make_dedupe_key(company.canonical_name, title, location),
        status="new",
        discovered_at=datetime.now(UTC),
        company_apply_url=url,
        source_import_key=f"reg|{make_uuid()}",
    )
    session.add(job)
    session.commit()
    return job


def test_gpa_conflict_detected_and_omitted(session: Session, tmp_path: Path) -> None:
    root = build_conflicting_gpa_corpus(tmp_path / "resumes")
    summary = scan_resume_corpus(session, path=root, rebuild=True)
    session.commit()
    assert summary.conflict_warnings >= 1
    gpas = {
        f.gpa_value
        for f in session.query(ResumeFact).filter(ResumeFact.gpa_value.isnot(None))
    }
    assert gpas >= {"3.5", "3.6", "3.7"}
    assert (
        session.query(ResumeFact)
        .filter(ResumeFact.gpa_value.isnot(None), ResumeFact.usable.is_(True))
        .count()
        == 0
    )
    job = _job(
        session,
        title="Software Engineer",
        description="Backend Python API Docker Kubernetes software engineer " * 20,
    )
    strategy = build_resume_strategy(session, job, deterministic_only=True)
    assert strategy.document is not None
    assert all(ed.gpa is None for ed in strategy.document.education)
    docx = render_docx(strategy.document, tmp_path / "out.docx")
    text = extract_docx_text(docx)
    assert "GPA" not in text
    assert "3.5" not in text and "3.6" not in text and "3.7" not in text


def test_skills_deduped_once(session: Session, tmp_path: Path) -> None:
    root = build_standard_corpus(tmp_path / "resumes")
    scan_resume_corpus(session, path=root, rebuild=True)
    session.commit()
    facts = session.query(ResumeFact).filter(ResumeFact.section == "skills").all()
    lines = build_deduped_skill_lines(facts)
    tokens: list[str] = []
    for line in lines:
        tokens.extend(normalize_text(t) for t in line.split(":", 1)[1].split(","))
    assert len(tokens) == len(set(tokens))


def test_near_duplicate_azure_variants_collapsed() -> None:
    a = (
        "Built staging and production environments across AWS, GCP, and Azure "
        "with Terraform and Kubernetes through Jenkins."
    )
    b = (
        "Built multiple staging and production environments in AWS, GCP and Azure. "
        "Provisioned with Terraform and deployed to Kubernetes using Jenkins."
    )
    assert _is_near_duplicate(b, [a]) is True


def test_orphan_experience_bullet_excluded(session: Session, tmp_path: Path) -> None:
    root = tmp_path / "resumes"
    write_resume_docx(
        root / "Current_Resume_Dev.docx",
        sections={
            "Skills": ["Languages: Python"],
            "Experience": [
                # orphan bullet with no company/title header first
                "Reduced Azure costs by 40% somehow without a company header.",
                "Acme Corp | Software Engineer | Jan 2022 – Present",
                "Built REST APIs in Python.",
            ],
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
            ],
        },
    )
    scan_resume_corpus(session, path=root, rebuild=True)
    session.commit()
    orphans = (
        session.query(ResumeFact).filter(ResumeFact.fact_kind == "orphan_bullet").all()
    )
    assert orphans
    assert all(f.usable is False for f in orphans)
    doc = build_resume_document(session, track="dev", keywords=["python", "api"])
    blob = " ".join(b.text for e in doc.experience for b in e.bullets)
    assert "without a company header" not in blob


def test_experience_requires_company_title_dates() -> None:
    doc = ResumeDocument(
        name_line="Alex Michael Manjakunnel",
        contact_line="a@b.com",
        summary=None,
        skill_lines=["Languages: Python"],
        experience=[
            ExperienceEntry(
                company="",
                title="Engineer",
                dates="2022",
                bullets=[
                    BulletRef(text="Did work", fact_id="1", normalized="did work")
                ],
            )
        ],
        education=[
            EducationEntry(
                institution="PES University",
                degree="Bachelor of Technology",
                parent_key="pes",
            )
        ],
    )
    errors = validate_document_structure(doc)
    assert any("company" in e.lower() or "incomplete" in e.lower() for e in errors)


def test_project_requires_name() -> None:
    doc = ResumeDocument(
        name_line="Alex Michael Manjakunnel",
        contact_line=None,
        summary=None,
        skill_lines=["Languages: Python"],
        experience=[
            ExperienceEntry(
                company="Acme",
                title="Engineer",
                dates="2022",
                bullets=[
                    BulletRef(text="Built APIs", fact_id="1", normalized="built apis")
                ],
                parent_key="acme",
            )
        ],
        projects=[
            ProjectEntry(
                name="",
                bullets=[
                    BulletRef(text="Did TacoDB work", fact_id="2", normalized="tacodb")
                ],
            )
        ],
        education=[
            EducationEntry(
                institution="PES University",
                degree="Bachelor of Technology",
                parent_key="pes",
            )
        ],
    )
    errors = validate_document_structure(doc)
    assert any("project" in e.lower() for e in errors)


def test_excluded_clearance_job_blocked(session: Session, tmp_path: Path) -> None:
    root = build_standard_corpus(tmp_path / "resumes")
    scan_resume_corpus(session, path=root, rebuild=True)
    session.commit()
    job = _job(
        session,
        title="Senior Infrastructure Engineer",
        description=(
            "Active TS/SCI clearance required. U.S. citizenship required. "
            "Kubernetes AWS platform engineering. " * 10
        ),
    )
    assert generation_preflight(session, job) == ELIGIBILITY_BLOCK_MSG
    with patch(
        "apps.resume_generation.services.generation_service.convert_docx_to_pdf"
    ) as conv:
        result = generate_resume_for_job(
            session, job.id, deterministic_only=True, copy_to_desktop=True
        )
        session.commit()
        assert result.status == "blocked"
        assert result.error == ELIGIBILITY_BLOCK_MSG
        assert result.desktop_pdf_path is None
        conv.assert_not_called()
        gen = session.get(ResumeGeneration, result.generation_id)
        assert gen is not None
        assert gen.status == "blocked"


def test_company_specific_secret_claim_not_auto_used(
    session: Session, tmp_path: Path
) -> None:
    root = build_standard_corpus(tmp_path / "resumes")
    scan_resume_corpus(session, path=root, rebuild=True)
    session.commit()
    job = _job(
        session,
        title="Software Engineer",
        description="Backend Python API software engineer Docker " * 20,
    )
    strategy = build_resume_strategy(session, job, deterministic_only=True)
    assert strategy.document is not None
    blob = " ".join(
        b.text for e in strategy.document.experience for b in e.bullets
    ).lower()
    assert "secret product" not in blob
    assert "raytheon" not in blob
    assert "accuweather" not in blob


def test_rendered_docx_single_skills_and_parent_headers(tmp_path: Path) -> None:
    doc = ResumeDocument(
        name_line="Alex Michael Manjakunnel",
        contact_line="a@example.com",
        summary=None,
        skill_lines=["Languages: Python, Java", "Cloud & DevOps: Docker, AWS"],
        experience=[
            ExperienceEntry(
                company="Sprinklr",
                title="Associate Cloud Engineer",
                dates="Jul 2023 – Apr 2025",
                bullets=[
                    BulletRef(
                        text="Deployed apps to Kubernetes on AWS.",
                        fact_id="1",
                        normalized="deployed apps to kubernetes on aws.",
                    )
                ],
                parent_key="spr",
            )
        ],
        projects=[
            ProjectEntry(
                name="TacoDB",
                bullets=[
                    BulletRef(
                        text="Implemented B+-Tree indexing.",
                        fact_id="2",
                        normalized="implemented b+-tree indexing.",
                    )
                ],
                parent_key="taco",
            )
        ],
        education=[
            EducationEntry(
                institution="PES University",
                degree="Bachelor of Technology",
                major="Computer Science and Engineering",
                graduation_date="Jun 2023",
                parent_key="pes",
            )
        ],
    )
    path = render_docx(doc, tmp_path / "r.docx")
    text = extract_docx_text(path)
    assert text.upper().count("SKILLS") == 1
    assert text.upper().count("EXPERIENCE") == 1
    assert text.upper().count("PROJECTS") == 1
    assert text.upper().count("EDUCATION") == 1
    assert "Sprinklr" in text
    assert "Associate Cloud Engineer" in text
    assert "TacoDB" in text
    # bullet must appear after company header
    assert text.index("Sprinklr") < text.index("Deployed apps")
    assert text.index("TacoDB") < text.index("B+-Tree")


def test_validate_artifacts_rejects_duplicate_sections(tmp_path: Path) -> None:
    # Build a broken docx manually with two SKILLS headings
    from docx import Document

    bad = Document()
    bad.add_paragraph("Alex Michael Manjakunnel")
    bad.add_paragraph("SKILLS")
    bad.add_paragraph("Python")
    bad.add_paragraph("SKILLS")
    bad.add_paragraph("Java")
    bad.add_paragraph("EXPERIENCE")
    bad.add_paragraph("Acme | Eng | 2022")
    bad.add_paragraph("Did work.")
    bad.add_paragraph("EDUCATION")
    bad.add_paragraph("PES | B.Tech")
    path = tmp_path / "bad.docx"
    bad.save(str(path))
    pdf = tmp_path / "bad.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake empty")
    report = validate_artifacts(docx_path=path, pdf_path=pdf)
    assert report.ok is False
    assert any("SKILLS" in m for m in report.messages)
