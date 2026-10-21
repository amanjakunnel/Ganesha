"""Helpers to build sanitized DOCX resume fixtures for tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from sqlalchemy.orm import Session

from packages.core.db import SessionLocal, get_engine
from packages.core.domain.models import Base


@pytest.fixture()
def resume_db_session() -> Session:
    """Fresh schema for resume-generation tests (sqlite)."""
    engine = get_engine()
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def write_resume_docx(path: Path, *, sections: dict[str, list[str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("Alex Michael Manjakunnel")
    doc.add_paragraph("alex@example.com | linkedin.com/in/example | github.com/example")
    for heading, bullets in sections.items():
        doc.add_paragraph(heading.upper())
        for bullet in bullets:
            doc.add_paragraph(bullet)
    doc.save(str(path))
    return path


def build_standard_corpus(root: Path) -> Path:
    """Create a small multi-track corpus with proper parent headers."""
    write_resume_docx(
        root / "Current_Resume_Dev.docx",
        sections={
            "Skills": [
                "Languages: Python, Java, JavaScript, SQL",
                "Cloud & DevOps: Docker, Git, Jenkins",
            ],
            "Experience": [
                "Acme Corp | Software Engineer | Jan 2022 – Present",
                "Built REST APIs in Python serving internal tooling.",
                "Improved CI/CD pipelines with Jenkins and Git.",
            ],
            "Projects": [
                "TacoDB | C++, SQL",
                "Implemented a C++ disk-based relational DBMS with buffer management.",
                "Added persistent B+-Tree indexing and query optimization.",
            ],
            "Education": [
                "The Pennsylvania State University | Master of Science | Aug 2025 – May 2027",
                "Major in Computer Science and Engineering",
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.5/4.0",
            ],
        },
    )
    write_resume_docx(
        root / "Current_Resume_Cloud.docx",
        sections={
            "Skills": [
                "Languages: Python, Shell",
                "Cloud & DevOps: AWS, Kubernetes, Linux, Terraform, Docker",
            ],
            "Experience": [
                "CloudCo | Cloud Engineer Intern | Jun 2023 – Aug 2023",
                "Operated Linux services on AWS with Kubernetes deployments.",
                "Automated release engineering scripts for containerized apps.",
            ],
            "Projects": [
                "InfraLab | Terraform, AWS",
                "Provisioned reproducible cloud environments with Terraform.",
            ],
            "Education": [
                "The Pennsylvania State University | Master of Science | Aug 2025 – May 2027",
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.5/4.0",
            ],
        },
    )
    write_resume_docx(
        root / "Current_Resume_ML.docx",
        sections={
            "Skills": [
                "Languages: Python",
                "Data & ML: PyTorch, scikit-learn, SQL",
            ],
            "Experience": [
                "ModelLabs | ML Intern | May 2022 – Aug 2022",
                "Trained classification models in PyTorch for research prototypes.",
            ],
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.5/4.0",
            ],
        },
    )
    write_resume_docx(
        root / "Current_Resume.docx",
        sections={
            "Skills": ["Languages: Python, Java, SQL"],
            "Experience": [
                "Acme Corp | Software Engineer | Jan 2022 – Present",
                "Built REST APIs in Python serving internal tooling.",
            ],
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.5/4.0",
            ],
        },
    )
    write_resume_docx(
        root / "Company Specific" / "Acme_Tailored.docx",
        sections={
            "Skills": ["Languages: Python, Java"],
            "Experience": [
                "Acme Corp | Software Engineer | Jan 2022 – Present",
                "Delivered features exclusively for Acme secret product X (unverified claim).",
            ],
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.5/4.0",
            ],
        },
    )
    (root / "ignore_me.pdf").write_bytes(b"%PDF-1.4 fake")
    (root / "~$Current_Resume_Dev.docx").write_bytes(b"lock")
    (root / "notes.txt").write_text("not a resume", encoding="utf-8")
    return root


def build_conflicting_gpa_corpus(root: Path) -> Path:
    build_standard_corpus(root)
    write_resume_docx(
        root / "GPA_Conflict_36.docx",
        sections={
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.6/4.0",
            ],
        },
    )
    write_resume_docx(
        root / "GPA_Conflict_37.docx",
        sections={
            "Education": [
                "PES University | Bachelor of Technology | Aug 2019 – Jun 2023",
                "Major in Computer Science and Engineering GPA: 3.7/4.0",
            ],
        },
    )
    return root
