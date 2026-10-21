from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from apps.resume_generation.paths import TEMPLATE_VERSION
from apps.resume_generation.services.resume_composer_service import ComposedResume


def render_docx(composed: ComposedResume, output_path: Path) -> Path:
    """Render a single-column ATS-safe DOCX from composed verified content."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(composed.name_line)
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Calibri"
    name.paragraph_format.space_after = Pt(2)

    if composed.contact_line:
        contact = doc.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = contact.add_run(composed.contact_line)
        c_run.font.size = Pt(9)
        c_run.font.name = "Calibri"
        contact.paragraph_format.space_after = Pt(8)

    def add_heading(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = "Calibri"
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)

    def add_bullets(items: list[str]) -> None:
        for item in items:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(item)
            r.font.size = Pt(10)
            r.font.name = "Calibri"
            p.paragraph_format.space_after = Pt(1)

    if composed.summary:
        add_heading("Summary")
        p = doc.add_paragraph()
        r = p.add_run(composed.summary)
        r.font.size = Pt(10)
        r.font.name = "Calibri"

    if composed.skills:
        add_heading("Skills")
        add_bullets(composed.skills)

    if composed.experience:
        add_heading("Experience")
        add_bullets(composed.experience)

    if composed.projects:
        add_heading("Projects")
        add_bullets(composed.projects)

    if composed.education:
        add_heading("Education")
        add_bullets(composed.education)

    # Document property for template version (not body text — ATS validators strip metadata checks)
    doc.core_properties.comments = f"template:{TEMPLATE_VERSION}"
    doc.save(str(output_path))
    return output_path


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text and p.text.strip())
